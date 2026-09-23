r"""Build the Word article file the submission portal requires.

    .\.venv\Scripts\python.exe scripts\06_docx.py

The portal rejects PDF for the Article File, so this renders results/letter.tex
into results/letter.docx rather than retyping it -- the .tex stays the single
source and the two cannot drift.

Formatting follows the journal: Times New Roman 12 pt, 1 inch margins, title
page single-spaced, remaining text double-spaced, page numbering beginning
with the title page, superscript reference numerals, and the reference list
double-spaced on its own page with journal titles italicised.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "results" / "letter.tex"
OUT = ROOT / "results" / "letter.docx"

CITE_OPEN, CITE_CLOSE = "\x01", "\x02"      # sentinels for superscript runs
ACCENTS = {r"\~a": "\u00e3", r"\~o": "\u00f5", r"\'a": "\u00e1",
           r"\'e": "\u00e9", r"\'i": "\u00ed", r"\'o": "\u00f3",
           r"\^a": "\u00e2", r"\c{c}": "\u00e7"}
# Escaped specials must survive comment stripping: a bare regex on "%" would
# eat "\%" and everything after it on the line.
GUARD = {r"\%": "\x03", r"\&": "\x04", r"\_": "\x05", r"\#": "\x06"}
UNGUARD = {"\x03": "%", "\x04": "&", "\x05": "_", "\x06": "#"}


def detex(t: str) -> str:
    """LaTeX source to plain text, preserving the characters that matter."""
    for k, v in GUARD.items():
        t = t.replace(k, v)
    t = "\n".join(ln for ln in t.split("\n")
                  if not ln.lstrip().startswith("%"))
    t = re.sub(r"%.*", "", t)                        # trailing comments
    for k, v in ACCENTS.items():
        t = t.replace(k, v)
    t = t.replace("``", "\u201c").replace("''", "\u201d")
    t = t.replace("---", "\u2014").replace("--", "\u2013")
    t = re.sub(r"\\includegraphics\[[^\]]*\]\{[^}]*\}", "", t)
    t = re.sub(r"\\makebox\[[^\]]*\]\[[^\]]*\]", "", t)
    # Separate adjacent commands first. Otherwise \noindent\textbf{Figure 1}
    # unwraps to \noindentFigure 1 and the bare-command strip below swallows
    # "Figure" as part of the command name.
    t = re.sub(r"(\\[a-zA-Z]+)(?=\\)", r"\1 ", t)
    for _ in range(3):                               # nested \textbf{\textit{}}
        t = re.sub(r"\\[a-zA-Z]+\*?\{([^{}]*)\}", r"\1", t)
    t = re.sub(r"\\[a-zA-Z]+\*?", " ", t)            # bare commands
    t = t.replace("$", "").replace("{", "").replace("}", "").replace("\\", "")
    for k, v in UNGUARD.items():
        t = t.replace(k, v)
    return re.sub(r"\s+", " ", t).strip()


def page_numbers(section) -> None:
    """A PAGE field in the footer, so numbering starts at the title page."""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)


def add(doc, text, *, double: bool, indent: float = 0.0,
        bold: bool = False, italics=()) -> None:
    """One paragraph; CITE sentinels become superscript runs."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = (WD_LINE_SPACING.DOUBLE if double
                            else WD_LINE_SPACING.SINGLE)
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(indent)

    for piece in re.split(f"({CITE_OPEN}[^{CITE_CLOSE}]*{CITE_CLOSE})", text):
        if not piece:
            continue
        if piece.startswith(CITE_OPEN):
            r = p.add_run(piece[1:-1])
            r.font.superscript = True
            continue
        cursor = 0
        for span in italics:
            i = piece.find(span, cursor)
            if i < 0:
                continue
            if i > cursor:
                p.add_run(piece[cursor:i]).bold = bold
            r = p.add_run(span)
            r.italic, r.bold = True, bold
            cursor = i + len(span)
        p.add_run(piece[cursor:]).bold = bold


def main() -> None:
    src = SRC.read_text(encoding="utf-8")
    keys = re.findall(r"\\bibitem\{([^}]*)\}", src)
    num = {k: i + 1 for i, k in enumerate(keys)}

    def cited(t: str) -> str:
        return re.sub(
            r"\\cite\{([^}]*)\}",
            lambda m: CITE_OPEN + ",".join(
                str(num[k.strip()]) for k in m.group(1).split(",")
            ) + CITE_CLOSE, t)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    page_numbers(sec)

    # ---- title page, single-spaced ---------------------------------------
    title = re.search(r"\\bfseries\\large (.*?)\}\s*\n", src, re.S).group(1)
    add(doc, detex(title), double=False, bold=True)
    add(doc, "", double=False)
    # One author, one affiliation: the superscript marker is noise, and
    # "$^{a}$" would otherwise survive detex as a literal "^a".
    author = re.search(r"\\noindent (Khanh[^\n]*)", src)
    if author:
        add(doc, detex(author.group(1).replace("$^{a}$", "")), double=False)
    affil = re.search(r"\\noindent \$\^\{a\}\$(.*?)\n\s*\n", src, re.S)
    if affil:
        add(doc, detex(affil.group(1)), double=False)

    # ---- body, double-spaced, from its own page --------------------------
    doc.add_page_break()
    start = src.index("RESEARCH LETTER")
    start = src.index("\n", start) + 1
    body = src[start:src.index(r"\begin{thebibliography}")]
    body = re.sub(r"\\clearpage|\\newpage|\\nolinenumbers|\\vfill", "", body)
    for para in re.split(r"\n\s*\n", body):
        text = detex(cited(para))
        if text:
            add(doc, text, double=True, indent=0.5)

    # ---- references, double-spaced, own page -----------------------------
    doc.add_page_break()
    add(doc, "References", double=True, bold=True)
    refs = src[src.index(r"\begin{thebibliography}"):
               src.index(r"\end{thebibliography}")]
    for i, chunk in enumerate(re.split(r"\\bibitem\{[^}]*\}", refs)[1:], 1):
        journals = [detex(j) for j in re.findall(r"\\textit\{([^}]*)\}", chunk)]
        add(doc, f"{i}. " + detex(chunk), double=True, italics=journals)

    # ---- figure legend, own page -----------------------------------------
    m = re.search(r"\\noindent\\textbf\{Figure 1\..*", src, re.S)
    if m:
        leg = m.group(0)[:m.group(0).index(r"\end{document}")]
        text = detex(leg)
        if text:
            doc.add_page_break()
            add(doc, text, double=True)

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  ({len(keys)} references)")


if __name__ == "__main__":
    main()
